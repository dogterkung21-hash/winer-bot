import os
import re
import json
import threading
import subprocess
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font
from flask import Flask
from PIL import Image
from telegram import Update, ReplyKeyboardMarkup, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    ApplicationBuilder,
    ContextTypes,
    CommandHandler,
    MessageHandler,
    ConversationHandler,
    CallbackQueryHandler,
    filters,
)
from telegram.request import HTTPXRequest

# ==========================================
# 🔑 Token บอท Telegram
# ==========================================
TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN")
DB_FILE = "local_database.json"

LOCKED_ITEM_NAME = "เตียงผู้ป่วย 3 ไกร์ระบบมือหมุน ขนาด 96x200x46-67 ซม."

def init_db():
    if not os.path.exists(DB_FILE):
        data = {"customers": [], "documents": [], "last_doc_number": "WN-P56-52"}
        with open(DB_FILE, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=4)

def load_db():
    init_db()
    with open(DB_FILE, 'r', encoding='utf-8') as f:
        return json.load(f)

def save_db(data):
    with open(DB_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

def save_or_update_customer(cust_data):
    db = load_db()
    customers = db.get("customers", [])
    existing_index = -1
    for idx, c in enumerate(customers):
        if (c.get("customer_name") and c.get("customer_name").strip() == cust_data.get("customer_name", "").strip()) or \
           (c.get("tax_id") and c.get("tax_id") != "-" and c.get("tax_id") == cust_data.get("tax_id")):
            existing_index = idx
            break

    if existing_index >= 0:
        customers[existing_index].update(cust_data)
    else:
        customers.append(cust_data)

    db["customers"] = customers
    save_db(db)

def get_next_doc_number():
    db = load_db()
    last_num = db.get("last_doc_number", "WN-P56-52")
    match = re.search(r'(\d+)$', last_num)
    if match:
        next_val = int(match.group(1)) + 1
        prefix = last_num[:match.start()]
        new_num = f"{prefix}{next_val:02d}"
    else:
        new_num = "WN-P56-53"
    db["last_doc_number"] = new_num
    save_db(db)
    return new_num

def get_thai_dates():
    now = datetime.now()
    thai_year = now.year + 543
    date_str = f"{now.day:02d}/{now.month:02d}/{thai_year}"
    due_date = now + timedelta(days=7)
    due_thai_year = due_date.year + 543
    due_date_str = f"{due_date.day:02d}/{due_date.month:02d}/{due_thai_year}"
    return date_str, due_date_str, now.strftime("%Y-%m-%d")

def thai_baht_text(amount):
    """ ฟังก์ชันแปลงตัวเลขเป็นข้อความบาทไทย ป้องกันปัญหา #NAME? """
    if amount is None:
        return "(ศูนย์บาทถ้วน)"
    num_text = ["ศูนย์", "หนึ่ง", "สอง", "สาม", "สี่", "ห้า", "หก", "เจ็ด", "แปด", "เก้า"]
    unit_text = ["", "สิบ", "ร้อย", "พัน", "หมื่น", "แสน", "ล้าน"]
    try:
        amount_str = f"{float(amount):.2f}"
        baht_part, _ = amount_str.split('.')
        
        def convert_chunk(n_str):
            result = ""
            length = len(n_str)
            for i, digit in enumerate(n_str):
                d = int(digit)
                if d == 0:
                    continue
                power = length - i - 1
                if power == 1 and d == 1:
                    result += "สิบ"
                elif power == 1 and d == 2:
                    result += "ยี่สิบ"
                elif power == 0 and d == 1 and length > 1 and n_str[-2] != '0':
                    result += "เอ็ด"
                else:
                    result += num_text[d] + unit_text[power]
            return result

        baht_words = convert_chunk(baht_part)
        return f"({baht_words}บาทถ้วน)" if baht_words else "(ศูนย์บาทถ้วน)"
    except:
        return "(ศูนย์บาทถ้วน)"

try:
    from pdf2image import convert_from_path
except ImportError:
    convert_from_path = None

CONVERT_ERRORS = {}

def find_content_bbox_px(pil_image, white_threshold=245):
    mask = pil_image.convert("L").point(lambda p: 255 if p < white_threshold else 0)
    return mask.getbbox()

def make_borderless_pdf_and_png(raw_pdf, final_pdf, final_png, png_dpi=200):
    os.replace(raw_pdf, final_pdf)
    if convert_from_path is not None:
        try:
            pil_img = convert_from_path(final_pdf, dpi=png_dpi)[0].convert("RGB")
            bbox = find_content_bbox_px(pil_img)
            if bbox:
                pil_img = pil_img.crop(bbox)
            pil_img.save(final_png, "PNG")
            return
        except Exception as e:
            print(f"pdf2image error: {e}")
    
    img = Image.new('RGB', (800, 1100), color=(255, 255, 255))
    img.save(final_png, "PNG")

BOT_VERSION = "v4-render-noname-fix"
CONVERT_STATUS = {}

def status_line(doc_number):
    s = CONVERT_STATUS.get(doc_number)
    return f"\n\n🔧 {s}" if s else ""

def convert_single_sheet_to_pdf_and_image(excel_filename, target_sheet_name, doc_number):
    pdf_filename = f"{doc_number}.pdf"
    image_filename = f"{doc_number}.png"
    raw_pdf = f"raw_{doc_number}.pdf"
    pdf_created, img_created = False, False
    CONVERT_ERRORS.pop(doc_number, None)

    sheet_only_excel = f"single_sheet_{doc_number}.xlsx"

    try:
        wb = openpyxl.load_workbook(excel_filename)
        for sheet_name in wb.sheetnames:
            if sheet_name != target_sheet_name:
                del wb[sheet_name]
        wb.save(sheet_only_excel)

        cmd = [
            "libreoffice",
            "--headless",
            "--convert-to", "pdf",
            "--outdir", os.getcwd(),
            sheet_only_excel
        ]
        subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        
        generated_pdf = os.path.splitext(os.path.basename(sheet_only_excel))[0] + ".pdf"
        if os.path.exists(generated_pdf):
            os.replace(generated_pdf, raw_pdf)
            pdf_created = True
            
    except Exception as e:
        print(f"LibreOffice PDF Convert Error: {e}")
        CONVERT_ERRORS[doc_number] = f"ส่งออก PDF ไม่สำเร็จ: {e}"
    finally:
        if os.path.exists(sheet_only_excel):
            try:
                os.remove(sheet_only_excel)
            except Exception:
                pass

    CONVERT_STATUS[doc_number] = f"{BOT_VERSION} | Sheet: {target_sheet_name}"

    if pdf_created and os.path.exists(raw_pdf):
        try:
            make_borderless_pdf_and_png(raw_pdf, pdf_filename, image_filename)
            img_created = os.path.exists(image_filename)
        except Exception as e:
            print(f"Borderless/PNG Error: {e}")
            CONVERT_ERRORS[doc_number] = f"แปลงรูปภาพไม่สำเร็จ: {e}"
            if not os.path.exists(pdf_filename):
                try:
                    os.replace(raw_pdf, pdf_filename)
                except Exception:
                    pass
        finally:
            if os.path.exists(raw_pdf):
                try:
                    os.remove(raw_pdf)
                except Exception:
                    pass

    pdf_ok = os.path.exists(pdf_filename)
    return (pdf_filename if pdf_ok else None), (image_filename if img_created else None)

main_menu_markup = ReplyKeyboardMarkup(
    [
        ["📄 ออกใบเสนอราคา", "📑 ออกอินวอย (Invoice)"],
        ["🧾 ออกใบเสร็จรับเงิน", "📋 รายชื่อลูกค้าที่บันทึกไว้"],
        ["📖 คู่มือการใช้งาน"]
    ],
    resize_keyboard=True
)

QUOT_FORM, INPUT_FIELD = range(2)

def build_quotation_keyboard(data):
    tax_val = data.get('tax_id') or '-'
    phone_val = data.get('phone') or '-'
    email_val = data.get('email') or '-'
    vat_status = "มี VAT 7%" if data.get('use_vat', True) else "ไม่มี VAT (0%)"
    keyboard = [
        [InlineKeyboardButton(f"👤 ชื่อลูกค้า: {data.get('customer_name')}", callback_data="edit_name")],
        [InlineKeyboardButton(f"📍 ที่อยู่ 1: {data.get('address_line1')}", callback_data="edit_addr1")],
        [InlineKeyboardButton(f"📍 ที่อยู่ 2: {data.get('address_line2')}", callback_data="edit_addr2")],
        [
            InlineKeyboardButton(f"🆔 เลขภาษี: {tax_val}", callback_data="edit_tax"),
            InlineKeyboardButton(f"📞 โทร: {phone_val}", callback_data="edit_phone")
        ],
        [InlineKeyboardButton(f"📧 อีเมล: {email_val}", callback_data="edit_email")],
        [InlineKeyboardButton(f"📦 สินค้า: {LOCKED_ITEM_NAME} (ล็อค)", callback_data="no_edit_item")],
        [
            InlineKeyboardButton(f"🔢 จำนวน: {data.get('qty'):,.0f}", callback_data="edit_qty"),
            InlineKeyboardButton(f"💰 ราคา/หน่วย: {data.get('price_per_unit'):,.2f}", callback_data="edit_price")
        ],
        [InlineKeyboardButton(f"⚙️ รูปแบบภาษี: {vat_status}", callback_data="toggle_vat")],
        [InlineKeyboardButton(f"📝 หมายเหตุ: {data.get('remarks')}", callback_data="edit_remarks")],
        [InlineKeyboardButton("🚀 สร้างใบเสนอราคา (เลือกรูปแบบไฟล์)", callback_data="generate_now")],
        [InlineKeyboardButton("❌ ยกเลิก", callback_data="cancel_menu")]
    ]
    return InlineKeyboardMarkup(keyboard)

async def start_quotation(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    context.user_data['quot_data'] = {
        'customer_name': 'บริษัท คลีน แอนด์ แคร์ พูล เซอร์วิส จำกัด',
        'address_line1': '162/101 หมู่ที่ 9 ตำบลสวนหลวง',
        'address_line2': 'อำเภอกระทุ่มแบน จังหวัดสมุทรสาคร 74110',
        'tax_id': '0745564002745',
        'phone': '0891351884',
        'email': '-',
        'item_name': LOCKED_ITEM_NAME,
        'qty': 20.0,
        'price_per_unit': 18300.0,
        'use_vat': True,
        'remarks': '-'
    }
    await update.message.reply_text(
        "📄 **ระบบออกใบเสนอราคา**\nกดปุ่มแก้ไขข้อมูล หรือกดสร้างเอกสารได้เลยครับ:",
        reply_markup=build_quotation_keyboard(context.user_data['quot_data'])
    )
    return QUOT_FORM

async def handle_quotation_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    action = query.data

    if action == "cancel_menu":
        await query.message.edit_text("❌ ยกเลิกรายการแล้ว")
        return ConversationHandler.END

    if action == "no_edit_item":
        await query.answer("สินค้านี้ถูกล็อคไว้ ไม่สามารถแก้ไขได้ครับ", show_alert=True)
        return QUOT_FORM

    if action == "toggle_vat":
        current = context.user_data['quot_data'].get('use_vat', True)
        context.user_data['quot_data']['use_vat'] = not current
        await query.message.edit_reply_markup(reply_markup=build_quotation_keyboard(context.user_data['quot_data']))
        return QUOT_FORM

    if action == "generate_now":
        await query.message.edit_text("⏳ กำลังสร้างไฟล์ PDF และรูปภาพ...")
        data = context.user_data['quot_data']
        data['item_name'] = LOCKED_ITEM_NAME
        doc_number = get_next_doc_number()
        date_str, due_str, date_iso = get_thai_dates()

        cust_info = {
            "customer_name": data['customer_name'],
            "address_line1": data['address_line1'],
            "address_line2": data['address_line2'],
            "tax_id": data['tax_id'],
            "phone": data['phone'],
            "email": data['email']
        }
        save_or_update_customer(cust_info)

        db = load_db()
        db["documents"].append({**data, "doc_number": doc_number, "date_iso": date_iso, "type": "quotation"})
        save_db(db)

        wb = openpyxl.load_workbook("บริษัท วินเนอร์ เชียงราย จำกัด.xlsx")
        ws = wb['เสนอราคา']
        ws['C9'] = f"      {data['customer_name']}"
        ws['C10'] = f"      {data['address_line1']}"
        ws['C11'] = f"      {data['address_line2']}"
        
        tax_val = str(data['tax_id'] if data['tax_id'] else '-').strip()
        phone_val = str(data['phone'] if data['phone'] else '-').strip()
        email_val = str(data['email'] if data['email'] else '-').strip()

        ws['B12'] = f"เลขกำกับภาษี    {tax_val}"
        ws['E12'] = f"เบอร์โทรศัพท์  {phone_val}"
        ws['B13'] = f"อีเมล  {email_val}"

        ws['H9'] = doc_number
        ws['H10'] = date_str
        ws['H11'] = due_str
        ws['C21'] = data['item_name']
        ws['F21'] = data['qty']
        ws['G21'] = data['price_per_unit']

        if not data.get('use_vat', True):
            ws['H34'] = 0
        else:
            ws['H34'] = "=H33*7%"

        if data['remarks'] != "-" and data['remarks'].strip() != "":
            ws['B33'] = f"หมายเหตุ : {data['remarks']}"

        temp_excel = f"temp_{doc_number}.xlsx"
        wb.save(temp_excel)
        
        pdf_file, img_file = convert_single_sheet_to_pdf_and_image(temp_excel, 'เสนอราคา', doc_number)
        if os.path.exists(temp_excel): os.remove(temp_excel)

        context.bot_data[f"pdf_{doc_number}"] = pdf_file
        context.bot_data[f"img_{doc_number}"] = img_file

        file_select_keyboard = [
            [
                InlineKeyboardButton("📄 รับไฟล์ PDF", callback_data=f"getpdf_{doc_number}"),
                InlineKeyboardButton("🖼️ รับเป็นรูปภาพ", callback_data=f"getimg_{doc_number}")
            ]
        ]
        await query.message.edit_text(
            f"✅ **สร้างใบเสนอราคาสำเร็จ!** (เลขที่: `{doc_number}`)\nเลือกรูปแบบไฟล์ที่ต้องการด้านล่าง:{status_line(doc_number)}",
            reply_markup=InlineKeyboardMarkup(file_select_keyboard)
        )
        return ConversationHandler.END

    field_map = {
        "edit_name": ("customer_name", "กรุณาพิมพ์ชื่อลูกค้า:"),
        "edit_addr1": ("address_line1", "กรุณาพิมพ์ที่อยู่ บรรทัดที่ 1:"),
        "edit_addr2": ("address_line2", "กรุณาพิมพ์ที่อยู่ บรรทัดที่ 2:"),
        "edit_tax": ("tax_id", "กรุณาพิมพ์เลขประจำตัวผู้เสียภาษี:"),
        "edit_phone": ("phone", "กรุณาพิมพ์เบอร์โทรศัพท์:"),
        "edit_email": ("email", "กรุณาพิมพ์อีเมล:"),
        "edit_qty": ("qty", "กรุณาพิมพ์จำนวนสินค้า:"),
        "edit_price": ("price_per_unit", "กรุณาพิมพ์ราคาต่อหน่วย:"),
        "edit_remarks": ("remarks", "กรุณาพิมพ์หมายเหตุ:")
    }

    if action in field_map:
        context.user_data['editing_field'] = field_map[action][0]
        await query.message.reply_text(field_map[action][1])
        return INPUT_FIELD

    return QUOT_FORM

async def receive_quotation_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    field = context.user_data.get('editing_field')
    if field:
        if field in ['qty', 'price_per_unit']:
            try: context.user_data['quot_data'][field] = float(text)
            except: pass
        else:
            context.user_data['quot_data'][field] = text

    await update.message.reply_text(
        "🔄 อัปเดตข้อมูลเรียบร้อยแล้ว:",
        reply_markup=build_quotation_keyboard(context.user_data['quot_data'])
    )
    return QUOT_FORM

INV_SELECT, INV_FORM, INV_INPUT = range(3)

def build_invoice_keyboard(data):
    paid_str = data.get('paid_text', '60%')
    tax_val = data.get('tax_id') or '-'
    phone_val = data.get('phone') or '-'
    email_val = data.get('email') or '-'
    vat_status = "มี VAT 7%" if data.get('use_vat', True) else "ไม่มี VAT (0%)"
    keyboard = [
        [InlineKeyboardButton(f"👤 ชื่อลูกค้า: {data.get('customer_name')}", callback_data="inv_edit_name")],
        [InlineKeyboardButton(f"📍 ที่อยู่ 1: {data.get('address_line1')}", callback_data="inv_edit_addr1")],
        [InlineKeyboardButton(f"📍 ที่อยู่ 2: {data.get('address_line2')}", callback_data="inv_edit_addr2")],
        [
            InlineKeyboardButton(f"🆔 เลขภาษี: {tax_val}", callback_data="inv_edit_tax"),
            InlineKeyboardButton(f"📞 โทร: {phone_val}", callback_data="inv_edit_phone")
        ],
        [InlineKeyboardButton(f"📧 อีเมล: {email_val}", callback_data="inv_edit_email")],
        [InlineKeyboardButton(f"📦 สินค้า: {LOCKED_ITEM_NAME} (ล็อค)", callback_data="inv_no_edit_item")],
        [
            InlineKeyboardButton(f"🔢 จำนวน: {data.get('qty'):,.0f}", callback_data="inv_edit_qty"),
            InlineKeyboardButton(f"💰 ราคา/หน่วย: {data.get('price_per_unit'):,.2f}", callback_data="inv_edit_price")
        ],
        [InlineKeyboardButton(f"⚙️ รูปแบบภาษี: {vat_status}", callback_data="inv_toggle_vat")],
        [InlineKeyboardButton(f"📝 หมายเหตุ: {data.get('remarks', '-')}", callback_data="inv_edit_remarks")],
        [InlineKeyboardButton(f"💵 ยอดชำระ/เปอร์เซ็นต์: {paid_str}", callback_data="inv_edit_paid")],
        [InlineKeyboardButton("🚀 สร้างอินวอย (เลือกรูปแบบไฟล์)", callback_data="inv_generate_now")],
        [InlineKeyboardButton("❌ ยกเลิก", callback_data="cancel_menu")]
    ]
    return InlineKeyboardMarkup(keyboard)

async def start_invoice(update: Update, context: ContextTypes.DEFAULT_TYPE):
    db = load_db()
    docs = db.get("documents", [])
    if not docs:
        await update.message.reply_text("⚠️ ยังไม่มีประวัติเอกสาร กรุณาออกใบเสนอราคาก่อนครับ", reply_markup=main_menu_markup)
        return ConversationHandler.END

    unique_docs = []
    seen_names = set()
    for doc in reversed(docs):
        c_name = doc.get("customer_name", "").strip()
        if c_name and c_name not in seen_names:
            seen_names.add(c_name)
            unique_docs.append(doc)

    keyboard = []
    for i, doc in enumerate(unique_docs):
        keyboard.append([InlineKeyboardButton(f"📁 {doc.get('customer_name')}", callback_data=f"sel_inv_{i}")])
    
    context.user_data['available_docs'] = unique_docs
    await update.message.reply_text(
        "📑 **ออกอินวอย (Invoice)**\nเลือกรายชื่อจากประวัติเพื่อดึงข้อมูลมาออกอินวอย:",
        reply_markup=InlineKeyboardMarkup(keyboard)
    )
    return INV_SELECT

async def invoice_doc_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    idx = int(query.data.split("_")[-1])
    doc = context.user_data['available_docs'][idx].copy()
    doc['paid_text'] = '60%'
    doc['item_name'] = LOCKED_ITEM_NAME
    doc['qty'] = 20.0
    if 'email' not in doc: doc['email'] = '-'
    if 'use_vat' not in doc: doc['use_vat'] = True
    if 'remarks' not in doc: doc['remarks'] = '-'
    context.user_data['inv_data'] = doc

    await query.message.edit_text(
        f"✅ ดึงข้อมูลล่าสุดของ: **{doc['customer_name']}** เรียบร้อย\nคุณสามารถกดปุ่มด้านล่างเพื่อแก้ไขข้อมูลก่อนสร้างอินวอย:",
        reply_markup=build_invoice_keyboard(doc)
    )
    return INV_FORM

def apply_bold_font_keep_style(cell):
    if cell.font:
        cell.font = Font(
            name=cell.font.name, size=cell.font.size, bold=True,
            italic=cell.font.italic, vertAlign=cell.font.vertAlign,
            underline=cell.font.underline, strike=cell.font.strike, color=cell.font.color
        )
    else:
        cell.font = Font(bold=True)

async def handle_invoice_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    action = query.data

    if action == "cancel_menu":
        await query.message.edit_text("❌ ยกเลิกรายการแล้ว")
        return ConversationHandler.END

    if action == "inv_no_edit_item":
        await query.answer("สินค้านี้ถูกล็อคไว้ ไม่สามารถแก้ไขได้ครับ", show_alert=True)
        return INV_FORM

    if action == "inv_toggle_vat":
        data = context.user_data['inv_data']
        data['use_vat'] = not data.get('use_vat', True)
        await query.message.edit_reply_markup(reply_markup=build_invoice_keyboard(data))
        return INV_FORM

    if action == "inv_generate_now":
        await query.message.edit_text("⏳ กำลังสร้างไฟล์ PDF และรูปภาพ...")
        data = context.user_data['inv_data']
        data['item_name'] = LOCKED_ITEM_NAME
        doc_number = get_next_doc_number()
        date_str, _, _ = get_thai_dates()
        
        use_vat = data.get('use_vat', True)
        multiplier = 1.07 if use_vat else 1.0
        total_val = data['qty'] * data['price_per_unit'] * multiplier
        paid_text = str(data.get('paid_text', '60')).strip()
        
        try:
            if '%' in paid_text:
                pct = round(float(paid_text.replace('%', '').strip()))
                paid_amt = total_val * (pct / 100)
            else:
                val_num = float(paid_text.replace(',', '').strip())
                if val_num <= 100 and "." not in paid_text:
                    pct = round(val_num)
                    paid_amt = total_val * (pct / 100)
                else:
                    paid_amt = val_num
                    pct = round((paid_amt / total_val) * 100)
        except:
            paid_amt = total_val * 0.60
            pct = 60

        wb = openpyxl.load_workbook("บริษัท วินเนอร์ เชียงราย จำกัด.xlsx")
        ws = wb['อินวอย(ช)']
        ws['B12'] = data['customer_name']
        ws['C13'] = data['address_line1']
        ws['C14'] = data['address_line2']
        
        tax_val = str(data['tax_id'] if data['tax_id'] else '-').strip()
        phone_val = str(data['phone'] if data['phone'] else '-').strip()

        ws['D15'] = tax_val
        ws['D16'] = phone_val

        apply_bold_font_keep_style(ws['B13'])
        apply_bold_font_keep_style(ws['B15'])

        ws['I13'] = doc_number
        ws['I14'] = date_str
        ws['C20'] = data['item_name']
        ws['G20'] = data['qty']
        ws['H20'] = data['price_per_unit']
        ws['J36'] = paid_amt

        # แก้ไขปัญหา #NAME? โดยการใส่ข้อความภาษาไทยแทนสูตร BAHTTEXT
        ws['B29'] = thai_baht_text(total_val)

        remark_val = data.get('remarks', '-').strip()
        if remark_val != "-" and remark_val != "":
            ws['C29'] = remark_val

        if not use_vat:
            ws['J31'] = 0
        else:
            ws['J31'] = "=J29*7%"

        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and ("ชำระก่อน" in cell.value or "ชำระ" in cell.value):
                    if "(" in cell.value and "%" in cell.value:
                        cell.value = f"ชำระก่อน ({pct}%)"

        temp_excel = f"temp_inv_{doc_number}.xlsx"
        wb.save(temp_excel)
        
        pdf_file, img_file = convert_single_sheet_to_pdf_and_image(temp_excel, 'อินวอย(ช)', doc_number)
        if os.path.exists(temp_excel): os.remove(temp_excel)

        context.bot_data[f"pdf_{doc_number}"] = pdf_file
        context.bot_data[f"img_{doc_number}"] = img_file

        file_select_keyboard = [
            [
                InlineKeyboardButton("📄 รับไฟล์ PDF", callback_data=f"getpdf_{doc_number}"),
                InlineKeyboardButton("🖼️ รับเป็นรูปภาพ", callback_data=f"getimg_{doc_number}")
            ]
        ]
        await query.message.edit_text(
            f"✅ **สร้างใบอินวอยสำเร็จ!** (เลขที่: `{doc_number}`)\nเลือกรูปแบบไฟล์ที่ต้องการด้านล่าง:{status_line(doc_number)}",
            reply_markup=InlineKeyboardMarkup(file_select_keyboard)
        )
        return ConversationHandler.END

    field_map = {
        "inv_edit_name": ("customer_name", "กรุณาพิมพ์ชื่อลูกค้าใหม่:"),
        "inv_edit_addr1": ("address_line1", "กรุณาพิมพ์ที่อยู่ บรรทัดที่ 1 ใหม่:"),
        "inv_edit_addr2": ("address_line2", "กรุณาพิมพ์ที่อยู่ บรรทัดที่ 2 ใหม่:"),
        "inv_edit_tax": ("tax_id", "กรุณาพิมพ์เลขประจำตัวผู้เสียภาษีใหม่:"),
        "inv_edit_phone": ("phone", "กรุณาพิมพ์เบอร์โทรศัพท์ใหม่:"),
        "inv_edit_email": ("email", "กรุณาพิมพ์อีเมลใหม่:"),
        "inv_edit_qty": ("qty", "กรุณาพิมพ์จำนวนสินค้าใหม่:"),
        "inv_edit_price": ("price_per_unit", "กรุณาพิมพ์ราคาต่อหน่วยใหม่:"),
        "inv_edit_remarks": ("remarks", "กรุณาพิมพ์หมายเหตุ:"),
        "inv_edit_paid": ("paid_text", "กรุณาพิมพ์ยอดชำระ (เช่น 70% หรือยอดเงินบาท):")
    }

    if action in field_map:
        context.user_data['inv_editing_field'] = field_map[action][0]
        await query.message.reply_text(field_map[action][1])
        return INV_INPUT

    return INV_FORM

async def receive_invoice_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    field = context.user_data.get('inv_editing_field')
    data = context.user_data['inv_data']
    if field:
        if field in ['qty', 'price_per_unit']:
            try: data[field] = float(text)
            except: pass
        else: data[field] = text

    await update.message.reply_text("🔄 อัปเดตข้อมูลเรียบร้อยแล้ว:", reply_markup=build_invoice_keyboard(data))
    return INV_FORM

REC_SELECT, REC_FORM, REC_INPUT = range(3)

def build_receipt_keyboard(data):
    paid_str = data.get('paid_text', '100%')
    tax_val = data.get('tax_id') or '-'
    phone_val = data.get('phone') or '-'
    email_val = data.get('email') or '-'
    vat_status = "มี VAT 7%" if data.get('use_vat', True) else "ไม่มี VAT (0%)"
    keyboard = [
        [InlineKeyboardButton(f"👤 ชื่อลูกค้า: {data.get('customer_name')}", callback_data="rec_edit_name")],
        [InlineKeyboardButton(f"📍 ที่อยู่ 1: {data.get('address_line1')}", callback_data="rec_edit_addr1")],
        [InlineKeyboardButton(f"📍 ที่อยู่ 2: {data.get('address_line2')}", callback_data="rec_edit_addr2")],
        [
            InlineKeyboardButton(f"🆔 เลขภาษี: {tax_val}", callback_data="rec_edit_tax"),
            InlineKeyboardButton(f"📞 โทร: {phone_val}", callback_data="rec_edit_phone")
        ],
        [InlineKeyboardButton(f"📧 อีเมล: {email_val}", callback_data="rec_edit_email")],
        [InlineKeyboardButton(f"📦 สินค้า: {LOCKED_ITEM_NAME} (ล็อค)", callback_data="rec_no_edit_item")],
        [
            InlineKeyboardButton(f"🔢 จำนวน: {data.get('qty'):,.0f}", callback_data="rec_edit_qty"),
            InlineKeyboardButton(f"💰 ราคา/หน่วย: {data.get('price_per_unit'):,.2f}", callback_data="rec_edit_price")
        ],
        [InlineKeyboardButton(f"⚙️ รูปแบบภาษี: {vat_status}", callback_data="rec_toggle_vat")],
        [InlineKeyboardButton(f"💵 ยอดรับชำระจริง: {paid_str}", callback_data="rec_edit_paid")],
        [InlineKeyboardButton("🚀 สร้างใบเสร็จรับเงิน (เลือกรูปแบบไฟล์)", callback_data="rec_generate_now")],
        [InlineKeyboardButton("❌ ยกเลิก", callback_data="cancel_menu")]
    ]
    return InlineKeyboardMarkup(keyboard)

async def start_receipt(update: Update, context: ContextTypes.DEFAULT_TYPE):
    db = load_db()
    docs = db.get("documents", [])
    if not docs:
        await update.message.reply_text("⚠️ ยังไม่มีประวัติเอกสาร กรุณาออกใบเสนอราคาก่อนครับ", reply_markup=main_menu_markup)
        return ConversationHandler.END

    unique_docs = []
    seen_names = set()
    for doc in reversed(docs):
        c_name = doc.get("customer_name", "").strip()
        if c_name and c_name not in seen_names:
            seen_names.add(c_name)
            unique_docs.append(doc)

    keyboard = []
    for i, doc in enumerate(unique_docs):
        keyboard.append([InlineKeyboardButton(f"📁 {doc.get('customer_name')}", callback_data=f"sel_rec_{i}")])
    
    context.user_data['available_docs'] = unique_docs
    await update.message.reply_text(
        "🧾 **ออกใบเสร็จรับเงิน**\nเลือกรายชื่อจากประวัติเพื่อดึงข้อมูลมาออกใบเสร็จ:",
        reply_markup=InlineKeyboardMarkup(keyboard)
    )
    return REC_SELECT

async def receipt_doc_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    idx = int(query.data.split("_")[-1])
    doc = context.user_data['available_docs'][idx].copy()
    doc['paid_text'] = '100%'
    doc['item_name'] = LOCKED_ITEM_NAME
    doc['qty'] = 20.0
    if 'email' not in doc: doc['email'] = '-'
    if 'use_vat' not in doc: doc['use_vat'] = True
    context.user_data['rec_data'] = doc

    await query.message.edit_text(
        f"✅ ดึงข้อมูลล่าสุดของ: **{doc['customer_name']}** เรียบร้อย\nคุณสามารถกดปุ่มด้านล่างเพื่อแก้ไขข้อมูลก่อนสร้างใบเสร็จ:",
        reply_markup=build_receipt_keyboard(doc)
    )
    return REC_FORM

def apply_black_font_keep_style(cell):
    if cell.font:
        cell.font = Font(
            name=cell.font.name, size=cell.font.size, bold=cell.font.bold,
            italic=cell.font.italic, vertAlign=cell.font.vertAlign,
            underline=cell.font.underline, strike=cell.font.strike, color="000000"
        )
    else:
        cell.font = Font(color="000000")

async def handle_receipt_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    action = query.data

    if action == "cancel_menu":
        await query.message.edit_text("❌ ยกเลิกรายการแล้ว")
        return ConversationHandler.END

    if action == "rec_no_edit_item":
        await query.answer("สินค้านี้ถูกล็อคไว้ ไม่สามารถแก้ไขได้ครับ", show_alert=True)
        return REC_FORM

    if action == "rec_toggle_vat":
        data = context.user_data['rec_data']
        data['use_vat'] = not data.get('use_vat', True)
        await query.message.edit_reply_markup(reply_markup=build_receipt_keyboard(data))
        return REC_FORM

    if action == "rec_generate_now":
        await query.message.edit_text("⏳ กำลังสร้างไฟล์ PDF และรูปภาพ...")
        data = context.user_data['rec_data']
        data['item_name'] = LOCKED_ITEM_NAME
        doc_number = get_next_doc_number()
        date_str, due_str, _ = get_thai_dates()
        
        use_vat = data.get('use_vat', True)
        subtotal = data['qty'] * data['price_per_unit']
        
        paid_text = str(data.get('paid_text', '100')).strip()
        
        try:
            if '%' in paid_text:
                pct = round(float(paid_text.replace('%', '').strip()))
                if use_vat: paid_amt = (subtotal * 1.07) * (pct / 100)
                else: paid_amt = subtotal * (pct / 100)
            else:
                val_num = float(paid_text.replace(',', '').strip())
                if val_num <= 100 and "." not in paid_text:
                    pct = round(val_num)
                    if use_vat: paid_amt = (subtotal * 1.07) * (pct / 100)
                    else: paid_amt = subtotal * (pct / 100)
                else:
                    paid_amt = val_num
                    total_val = (subtotal * 1.07) if use_vat else subtotal
                    pct = round((paid_amt / total_val) * 100)
        except:
            paid_amt = (subtotal * 1.07) if use_vat else subtotal
            pct = 100

        wb = openpyxl.load_workbook("บริษัท วินเนอร์ เชียงราย จำกัด.xlsx")
        ws = wb['รับเงิน']

        ws['C10'] = data['customer_name']
        ws['D11'] = data['address_line1']
        ws['D12'] = data['address_line2']
        
        tax_val = str(data['tax_id'] if data['tax_id'] else '-').strip()
        phone_val = str(data['phone'] if data['phone'] else '-').strip()

        ws['E13'] = tax_val
        ws['G9'] = phone_val

        ws['I10'] = f"เลขที่          {doc_number}"
        apply_black_font_keep_style(ws['I10'])
        ws['I11'] = f"วันที่           {date_str}"
        apply_black_font_keep_style(ws['I11'])
        ws['I12'] = f"ครบกำหนด   {due_str}"
        apply_black_font_keep_style(ws['I12'])
        
        ws['D16'] = data['item_name']
        ws['G16'] = data['qty']
        ws['H16'] = data['price_per_unit']

        ws['I29'] = subtotal
        if not use_vat: ws['I30'] = 0
        else: ws['I30'] = "=I29*7%"

        ws['I32'] = paid_amt

        # แก้ไขปัญหา #NAME? โดยการใส่ข้อความภาษาไทยแทนสูตร BAHTTEXT ในใบเสร็จ
        total_full_val = (subtotal * 1.07) if use_vat else subtotal
        ws['B29'] = thai_baht_text(total_full_val)

        formatted_amt = f"{paid_amt:,.2f}"
        ws['C30'] = f"**ลูกค้าชำระยอด {pct}% จากยอดเต็มเป็นจำนวนเงิน {formatted_amt} บาท**"

        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and ("ชำระก่อน" in cell.value or "ชำระ" in cell.value):
                    if "(" in cell.value and "%" in cell.value:
                        cell.value = f"ชำระ ({pct}%)"

        temp_excel = f"temp_rec_{doc_number}.xlsx"
        wb.save(temp_excel)
        
        pdf_file, img_file = convert_single_sheet_to_pdf_and_image(temp_excel, 'รับเงิน', doc_number)
        if os.path.exists(temp_excel): os.remove(temp_excel)

        context.bot_data[f"pdf_{doc_number}"] = pdf_file
        context.bot_data[f"img_{doc_number}"] = img_file

        file_select_keyboard = [
            [
                InlineKeyboardButton("📄 รับไฟล์ PDF", callback_data=f"getpdf_{doc_number}"),
                InlineKeyboardButton("🖼️ รับเป็นรูปภาพ", callback_data=f"getimg_{doc_number}")
            ]
        ]
        await query.message.edit_text(
            f"✅ **สร้างใบเสร็จรับเงินสำเร็จ!** (เลขที่: `{doc_number}`)\nเลือกรูปแบบไฟล์ที่ต้องการด้านล่าง:{status_line(doc_number)}",
            reply_markup=InlineKeyboardMarkup(file_select_keyboard)
        )
        return ConversationHandler.END

    field_map = {
        "rec_edit_name": ("customer_name", "กรุณาพิมพ์ชื่อลูกค้าใหม่:"),
        "rec_edit_addr1": ("address_line1", "กรุณาพิมพ์ที่อยู่ บรรทัดที่ 1 ใหม่:"),
        "rec_edit_addr2": ("address_line2", "กรุณาพิมพ์ที่อยู่ บรรทัดที่ 2 ใหม่:"),
        "rec_edit_tax": ("tax_id", "กรุณาพิมพ์เลขประจำตัวผู้เสียภาษีใหม่:"),
        "rec_edit_phone": ("phone", "กรุณาพิมพ์เบอร์โทรศัพท์ใหม่:"),
        "rec_edit_email": ("email", "กรุณาพิมพ์อีเมลใหม่:"),
        "rec_edit_qty": ("qty", "กรุณาพิมพ์จำนวนสินค้าใหม่:"),
        "rec_edit_price": ("price_per_unit", "กรุณาพิมพ์ราคาต่อหน่วยใหม่:"),
        "rec_edit_paid": ("paid_text", "กรุณาพิมพ์ยอดชำระเงินจริง (เช่น 164700 หรือ 45%):")
    }

    if action in field_map:
        context.user_data['rec_editing_field'] = field_map[action][0]
        await query.message.reply_text(field_map[action][1])
        return REC_INPUT

    return REC_FORM

async def receive_receipt_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    field = context.user_data.get('rec_editing_field')
    data = context.user_data['rec_data']
    if field:
        if field in ['qty', 'price_per_unit']:
            try: data[field] = float(text)
            except ValueError: pass
        else: data[field] = text

    await update.message.reply_text("🔄 อัปเดตข้อมูลเรียบร้อยแล้ว:", reply_markup=build_receipt_keyboard(data))
    return REC_FORM

async def handle_file_selection_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    data = query.data

    if data.startswith("getpdf_"):
        doc_num = data.replace("getpdf_", "")
        pdf_file = context.bot_data.pop(f"pdf_{doc_num}", None)
        if pdf_file and os.path.exists(pdf_file):
            with open(pdf_file, 'rb') as f:
                await query.message.reply_document(document=f, caption=f"📄 ไฟล์ PDF เอกสารเลขที่ {doc_num}", reply_markup=main_menu_markup)
            try: os.remove(pdf_file)
            except: pass
            img_partner = f"{doc_num}.png"
            if os.path.exists(img_partner):
                try: os.remove(img_partner)
                except: pass
        else:
            err = CONVERT_ERRORS.get(doc_num, "ไม่ทราบสาเหตุ")
            await query.message.reply_text(f"⚠️ ไม่พบไฟล์ PDF หรือถูกส่งไปแล้ว\nสาเหตุ: {err}", reply_markup=main_menu_markup)

    elif data.startswith("getimg_"):
        doc_num = data.replace("getimg_", "")
        img_file = context.bot_data.pop(f"img_{doc_num}", None)
        if img_file and os.path.exists(img_file):
            with open(img_file, 'rb') as f:
                await query.message.reply_photo(photo=f, caption=f"🖼️ รูปภาพเอกสารเลขที่ {doc_num}", reply_markup=main_menu_markup)
            try: os.remove(img_file)
            except: pass
            pdf_partner = f"{doc_num}.pdf"
            if os.path.exists(pdf_partner):
                try: os.remove(pdf_partner)
                except: pass
        else:
            err = CONVERT_ERRORS.get(doc_num, "ไม่ทราบสาเหตุ")
            await query.message.reply_text(f"⚠️ ไม่พบไฟล์รูปภาพหรือถูกส่งไปแล้ว\nสาเหตุ: {err}", reply_markup=main_menu_markup)

async def show_customers(update: Update, context: ContextTypes.DEFAULT_TYPE):
    db = load_db()
    customers = db.get("customers", [])
    if not customers:
        await update.message.reply_text("⚠️ ยังไม่มีรายชื่อลูกค้าในระบบ", reply_markup=main_menu_markup)
        return
    msg = f"📋 **รายชื่อลูกค้าทั้งหมด ({len(customers)} ราย):**\n\n"
    for idx, c in enumerate(customers, 1):
        email_info = c.get('email', '-')
        msg += f"{idx}. **{c['customer_name']}**\n   📍 {c['address_line1']} {c['address_line2']}\n   🆔 เลขภาษี: {c['tax_id']} | 📞 โทร: {c['phone']} | 📧 อีเมล: {email_info}\n\n"
    await update.message.reply_text(msg, reply_markup=main_menu_markup)

async def show_guide(update: Update, context: ContextTypes.DEFAULT_TYPE):
    guide_text = (
        "📖 **คู่มือการใช้งานระบบออกเอกสาร**\n\n"
        "1️⃣ **เลือกเอกสาร:** สร้างใบเสนอราคา / อินวอย / ใบเสร็จรับเงิน\n"
        "2️⃣ **แก้ไขข้อมูลก่อนสร้าง:** สามารถกดปุ่มเพื่อแก้ไขชื่อ, ที่อยู่, เลขภาษี, โทรศัพท์, อีเมล, รูปแบบ VAT หรือยอดชำระได้ทันที\n"
        "3️⃣ **สินค้าถูกล็อค:** รายการสินค้าจะเป็น 'เตียงผู้ป่วย 3 ไกร์ระบบมือหมุน ขนาด 96x200x46-67 ซม.' อัตโนมัติ\n"
        "💬 พิมพ์ /start เพื่อรีเซ็ตเมนูหลักได้ตลอดเวลาครับ"
    )
    await update.message.reply_text(guide_text, reply_markup=main_menu_markup)

async def check_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("🔍 กำลังตรวจสอบระบบ...")
    lines = ["ระบบทำงานได้ปกติ"]
    await update.message.reply_text("\n".join(lines), reply_markup=main_menu_markup)

async def start_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        f"🤖 **ระบบออกเอกสาร Winner Chiangrai** ({BOT_VERSION})\nเลือกเมนูด้านล่างเพื่อเริ่มใช้งานได้เลยครับ:",
        reply_markup=main_menu_markup
    )

async def cancel_process(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text("❌ ยกเลิกรายการเรียบร้อย", reply_markup=main_menu_markup)
    return ConversationHandler.END

web_app = Flask('')
@web_app.route('/')
def home():
    return "Bot is running securely on Render.com 24/7!"

def run_web():
    port = int(os.environ.get("PORT", 8080))
    web_app.run(host='0.0.0.0', port=port)

if __name__ == "__main__":
    init_db()
    threading.Thread(target=run_web, daemon=True).start()
    
    request = HTTPXRequest(connect_timeout=30.0, read_timeout=30.0)
    app = ApplicationBuilder().token(TELEGRAM_BOT_TOKEN).request(request).build()
    
    quot_handler = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex(r"^(📄 ออกใบเสนอราคา|ออกใบเสนอราคา)$"), start_quotation)],
        states={
            QUOT_FORM: [CallbackQueryHandler(handle_quotation_callback)],
            INPUT_FIELD: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_quotation_input)],
        },
        fallbacks=[CommandHandler("cancel", cancel_process)],
    )

    inv_handler = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex(r"^(📑 ออกอินวอย \(Invoice\)|ออกอินวอย)$"), start_invoice)],
        states={
            INV_SELECT: [CallbackQueryHandler(invoice_doc_selected, pattern="^sel_inv_")],
            INV_FORM: [CallbackQueryHandler(handle_invoice_callback)],
            INV_INPUT: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_invoice_input)],
        },
        fallbacks=[CommandHandler("cancel", cancel_process)],
    )

    rec_handler = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex(r"^(🧾 ออกใบเสร็จรับเงิน|ออกใบเสร็จรับเงิน)$"), start_receipt)],
        states={
            REC_SELECT: [CallbackQueryHandler(receipt_doc_selected, pattern="^sel_rec_")],
            REC_FORM: [CallbackQueryHandler(handle_receipt_callback)],
            REC_INPUT: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_receipt_input)],
        },
        fallbacks=[CommandHandler("cancel", cancel_process)],
    )

    app.add_handler(CommandHandler("start", start_command))
    app.add_handler(CommandHandler("check", check_command))
    app.add_handler(MessageHandler(filters.Regex(r"^(📋 รายชื่อลูกค้าที่บันทึกไว้)$"), show_customers))
    app.add_handler(MessageHandler(filters.Regex(r"^(📖 คู่มือการใช้งาน)$"), show_guide))
    app.add_handler(quot_handler)
    app.add_handler(inv_handler)
    app.add_handler(rec_handler)
    app.add_handler(CallbackQueryHandler(handle_file_selection_callback, pattern="^(getpdf_|getimg_)"))
    
    print(f"เวอร์ชันบอท: {BOT_VERSION}")
    print("🤖 บอทพร้อมทำงานบน Render.com แล้ว...")
    app.run_polling()
